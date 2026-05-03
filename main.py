# v2

import os
import json
import base64
import httpx
import io
import re
import asyncio
from pathlib import Path
from datetime import datetime
from typing import Any, Dict, List

from fastapi import FastAPI, File, UploadFile, HTTPException, Request
from fastapi.responses import HTMLResponse, JSONResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from dotenv import load_dotenv

load_dotenv()

# ── Paths — create folders BEFORE anything else ────────────────────────────
BASE_DIR = Path(__file__).parent
(BASE_DIR / "static").mkdir(parents=True, exist_ok=True)
(BASE_DIR / "uploads").mkdir(parents=True, exist_ok=True)
(BASE_DIR / "templates").mkdir(parents=True, exist_ok=True)

DATA_FILE  = BASE_DIR / "speaker_data.json"
GROQ_URL   = "https://api.groq.com/openai/v1/chat/completions"
GROQ_MODEL = "meta-llama/llama-4-scout-17b-16e-instruct"

# ── Rate-limit config ──────────────────────────────────────────────────────
MAX_RETRIES        = 5          # maximum retry attempts per request
BASE_BACKOFF       = 2.0        # seconds — doubled each retry (2, 4, 8, 16, 32)
INTER_REQUEST_DELAY = 1.5       # seconds to wait between successful requests

# ── App ────────────────────────────────────────────────────────────────────
app = FastAPI(title="Toastmasters Speech Tracker")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

app.mount("/static", StaticFiles(directory=BASE_DIR / "static"), name="static")
templates = Jinja2Templates(directory=BASE_DIR / "templates")


# ── Pydantic models ────────────────────────────────────────────────────────
class Speech(BaseModel):
    name: str
    level: int
    project: int
    meetingDate: str = ""

class MergeRequest(BaseModel):
    speeches: List[Speech]

class SpeakerDataModel(BaseModel):
    data: Dict[str, Any]


# ── Data helpers ───────────────────────────────────────────────────────────
def load_data() -> dict:
    if DATA_FILE.exists():
        with open(DATA_FILE, "r", encoding="utf-8") as f:
            return json.load(f)
    return {}


def save_data(data: dict):
    with open(DATA_FILE, "w", encoding="utf-8") as f:
        json.dump(data, f, indent=2, ensure_ascii=False)


def date_from_filename(filename: str) -> str:
    m = re.search(r"(\d{1,2})[a-zA-Z]*[_\s]+([A-Za-z]+)[_\s]+(\d{4})", filename or "")
    if m:
        return f"{m.group(1)} {m.group(2)} {m.group(3)}"
    return datetime.now().strftime("%d %b %Y")


def normalize_name(name: str) -> str:
    return " ".join(name.lower().strip().split())


def merge_speeches(existing: dict, speeches: list) -> dict:
    name_map = {normalize_name(k): k for k in existing}

    for s in speeches:
        raw_name = s.get("name", "").strip()
        if not raw_name:
            continue

        norm      = normalize_name(raw_name)
        canonical = name_map.get(norm, raw_name)
        if canonical not in existing:
            existing[canonical] = []
            name_map[norm] = canonical

        level     = int(s.get("level") or 0)
        project   = int(s.get("project") or 0)
        stype     = (s.get("speech_type") or "").strip() or None
        meet_date = s.get("meetingDate", "")

        def is_dup(x):
            x_stype = (x.get("speech_type") or "").strip() or None
            if stype and x_stype:
                return x_stype == stype and x.get("meetingDate") == meet_date
            return (
                (x.get("level") or 0) == level and
                (x.get("project") or 0) == project and
                x.get("meetingDate") == meet_date
            )

        if not any(is_dup(x) for x in existing[canonical]):
            existing[canonical].append({
                "level": level,
                "project": project,
                "speech_type": stype,
                "meetingDate": meet_date,
            })
    return existing


# ── Groq call with exponential-backoff retry ───────────────────────────────
async def call_groq_with_retry(
    api_key: str,
    payload: dict,
    max_retries: int = MAX_RETRIES,
    base_backoff: float = BASE_BACKOFF,
) -> dict:
    """
    POST to Groq. On HTTP 429 (rate limit) or 503 (overload) retry up to
    `max_retries` times with exponential back-off.

    Returns the parsed JSON response dict on success.
    Raises HTTPException on non-retryable errors or exhausted retries.
    """
    headers = {
        "Authorization": f"Bearer {api_key}",
        "Content-Type": "application/json",
    }

    for attempt in range(max_retries + 1):
        async with httpx.AsyncClient(timeout=90) as client:
            resp = await client.post(GROQ_URL, json=payload, headers=headers)

        # ── Success ────────────────────────────────────────────────────────
        if resp.status_code == 200:
            return resp.json()

        # ── Rate-limited or server overloaded → retry ──────────────────────
        if resp.status_code in (429, 503):
            if attempt == max_retries:
                raise HTTPException(
                    429,
                    f"Groq rate limit hit after {max_retries} retries. "
                    "Please wait a minute and try again, or reduce the number "
                    "of agendas uploaded at once."
                )

            # Honour Retry-After if provided, otherwise exponential back-off
            retry_after = resp.headers.get("retry-after") or resp.headers.get("x-ratelimit-reset-requests")
            if retry_after:
                try:
                    wait = float(retry_after)
                except ValueError:
                    wait = base_backoff * (2 ** attempt)
            else:
                wait = base_backoff * (2 ** attempt)   # 2, 4, 8, 16, 32 s

            print(
                f"[Groq] Rate limited (attempt {attempt + 1}/{max_retries}). "
                f"Retrying in {wait:.1f}s …"
            )
            await asyncio.sleep(wait)
            continue

        # ── Any other error → fail immediately ────────────────────────────
        try:
            msg = resp.json().get("error", {}).get("message", resp.text)
        except Exception:
            msg = resp.text
        raise HTTPException(resp.status_code, f"Groq error: {msg}")

    # Should never reach here
    raise HTTPException(500, "Unexpected error in Groq retry loop")


# ── Routes ─────────────────────────────────────────────────────────────────
@app.get("/", response_class=HTMLResponse)
async def home(request: Request):
    return templates.TemplateResponse("index.html", {"request": request})


@app.post("/api/extract")
async def extract(request: Request, file: UploadFile = File(...)):
    """Upload an agenda image → Groq extracts speakers → auto-saves to disk."""
    api_key = (
        request.headers.get("X-Groq-Api-Key")
        or os.getenv("GROQ_API_KEY", "")
    )
    if not api_key:
        raise HTTPException(
            400,
            "Groq API key missing. Pass X-Groq-Api-Key header or set GROQ_API_KEY in .env"
        )

    contents = await file.read()
    if not contents:
        raise HTTPException(400, "Empty file")

    mime = file.content_type or "image/jpeg"
    b64  = base64.b64encode(contents).decode()

    prompt = (
        "This is a Toastmasters meeting agenda image.\n\n"
        "Task 1: Read the meeting date from the agenda header "
        "(it appears near the top, e.g. Dec 7th 2025 or 7 Dec 2025). "
        "Format it as DD Mon YYYY (e.g. 07 Dec 2025).\n\n"
        "Task 2: Extract ONLY the speakers listed under the PREPARED SPEECH SECTION.\n"
        "For each speaker extract:\n"
        "- name: full name of the speaker\n"
        "- level: the Level number as integer (e.g. Level 1 -> 1, Level 4 -> 4). "
        "If the speech is DTM Completion use 5. If no level is mentioned use 0.\n"
        "- project: the Project number as integer (e.g. Project 3 -> 3). "
        "If no project number is mentioned use 0.\n"
        "- speech_type: a short label for special speeches. "
        "For normal speeches use null. For DTM use 'DTM'. For Better Speaker Series use 'BSS'. "
        "For Ice Breaker use 'Ice Breaker'.\n\n"
        "Return a single plain JSON object with no markdown and no extra text.\n"
        "Example: {\"meetingDate\": \"07 Dec 2025\", \"speeches\": ["
        "{\"name\": \"Full Name\", \"level\": 1, \"project\": 2, \"speech_type\": null}, "
        "{\"name\": \"Another Name\", \"level\": 0, \"project\": 0, \"speech_type\": \"DTM\"}]}\n\n"
        "If no prepared speech section exists return: "
        "{\"meetingDate\": \"07 Dec 2025\", \"speeches\": []}"
    )

    payload = {
        "model": GROQ_MODEL,
        "temperature": 0,
        "max_tokens": 600,
        "messages": [{
            "role": "user",
            "content": [
                {"type": "image_url", "image_url": {"url": f"data:{mime};base64,{b64}"}},
                {"type": "text", "text": prompt},
            ],
        }],
    }

    # ── Call Groq with automatic retry on rate limits ──────────────────────
    result = await call_groq_with_retry(api_key, payload)

    raw   = result["choices"][0]["message"]["content"].strip()
    clean = raw.replace("```json", "").replace("```", "").strip()

    try:
        parsed = json.loads(clean)
        if isinstance(parsed, dict):
            date     = parsed.get("meetingDate", datetime.now().strftime("%d %b %Y"))
            speeches = parsed.get("speeches", [])
        else:
            speeches = parsed
            date     = datetime.now().strftime("%d %b %Y")
    except json.JSONDecodeError:
        raise HTTPException(500, f"Could not parse model response: {raw}")

    for s in speeches:
        s["meetingDate"] = date

    # ── Auto-save extracted speeches ───────────────────────────────────────
    existing = load_data()
    updated  = merge_speeches(existing, speeches)
    save_data(updated)

    # Small courtesy delay so bulk uploads don't hammer the API
    await asyncio.sleep(INTER_REQUEST_DELAY)

    return {
        "speeches": speeches,
        "meetingDate": date,
        "filename": file.filename,
        "saved": True,
        "total_speakers": len(updated),
    }


@app.get("/api/speakers")
async def get_speakers():
    return load_data()


@app.post("/api/speakers/merge")
async def merge_speakers(body: MergeRequest):
    existing = load_data()
    updated  = merge_speeches(existing, [s.model_dump() for s in body.speeches])
    save_data(updated)
    return {"status": "saved", "total_speakers": len(updated)}


@app.delete("/api/speakers")
async def delete_speakers():
    save_data({})
    return {"status": "cleared"}


@app.get("/api/export")
async def export_excel():
    data = load_data()
    if not data:
        raise HTTPException(404, "No data to export. Extract some agendas first.")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Speech Tracker"

    hdr_font = Font(name="Arial", bold=True, color="FFFFFF", size=11)
    hdr_fill = PatternFill("solid", fgColor="8B2131")
    sub_fill = PatternFill("solid", fgColor="C8973A")
    sub_font = Font(name="Arial", bold=True, color="FFFFFF", size=10)
    center   = Alignment(horizontal="center", vertical="center")
    thin     = Border(
        left=Side(style="thin", color="DDDDDD"), right=Side(style="thin", color="DDDDDD"),
        top=Side(style="thin", color="DDDDDD"),  bottom=Side(style="thin", color="DDDDDD"),
    )
    alt_fill = PatternFill("solid", fgColor="FFF8EE")

    max_s = max((len(v) for v in data.values()), default=0)

    def hcell(r, c, val, font=None, fill=None):
        cell = ws.cell(r, c, val)
        cell.font      = font or hdr_font
        cell.fill      = fill or hdr_fill
        cell.alignment = center
        cell.border    = thin
        return cell

    hcell(1, 1, "Speaker Name")
    hcell(1, 2, "Total Speeches")
    col = 3
    for i in range(1, max_s + 1):
        ws.merge_cells(start_row=1, start_column=col, end_row=1, end_column=col + 1)
        hcell(1, col, f"Speech {i}", font=sub_font, fill=sub_fill)
        col += 2

    for c, label in [(1, "Name"), (2, "Count")]:
        cell = ws.cell(2, c, label)
        cell.font      = Font(name="Arial", bold=True, size=10)
        cell.alignment = center
        cell.border    = thin

    col = 3
    for _ in range(max_s):
        for label in ["Level / Project", "Meeting Date"]:
            cell = ws.cell(2, col, label)
            cell.font      = Font(name="Arial", bold=True, size=10)
            cell.alignment = center
            cell.border    = thin
            col += 1

    ws.row_dimensions[1].height = 22
    ws.row_dimensions[2].height = 18

    for ri, (name, speeches) in enumerate(sorted(data.items()), start=3):
        ss   = sorted(speeches, key=lambda x: (x.get("level") or 0, x.get("project") or 0))
        fill = alt_fill if ri % 2 == 0 else None

        def dcell(r, c, val):
            cell = ws.cell(r, c, val)
            cell.font      = Font(name="Arial", size=10)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border    = thin
            if fill:
                cell.fill = fill

        ws.cell(ri, 1, name).font = Font(name="Arial", size=10)
        ws.cell(ri, 1).alignment  = Alignment(horizontal="left", vertical="center")
        ws.cell(ri, 1).border     = thin
        if fill:
            ws.cell(ri, 1).fill = fill

        dcell(ri, 2, len(ss))
        col = 3
        for i in range(max_s):
            if i < len(ss):
                lvl   = ss[i].get('level') or 0
                prj   = ss[i].get('project') or 0
                stype = ss[i].get('speech_type')
                if stype:
                    dcell(ri, col, stype)
                elif lvl and prj:
                    dcell(ri, col, f"L{lvl} P{prj}")
                else:
                    dcell(ri, col, "—")
                dcell(ri, col + 1, ss[i].get("meetingDate", ""))
            else:
                dcell(ri, col, "—")
                dcell(ri, col + 1, "—")
            col += 2

        ws.row_dimensions[ri].height = 18

    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 14
    for i in range(max_s):
        ws.column_dimensions[openpyxl.utils.get_column_letter(3 + i * 2)].width     = 18
        ws.column_dimensions[openpyxl.utils.get_column_letter(3 + i * 2 + 1)].width = 16

    ws.freeze_panes = "A3"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=toastmasters_tracker.xlsx"},
    )

from mangum import Mangum

handler = Mangum(app)